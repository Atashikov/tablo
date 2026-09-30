import json,re
from openpyxl import load_workbook
U='/mnt/user-data/uploads/'
ws=load_workbook(U+'Материалы_курсов.xlsx')['Лист1']
cats=[{'id':'cat-mandatory','name':'Обязательные программы обучения'},
      {'id':'cat-professions','name':'Программы обучения по профессиям'},
      {'id':'cat-hse','name':'Обучение по программам HSE'}]
courses=[];cur=None;diff=[]
for r in range(1,ws.max_row+1):
    b=ws.cell(r,2).value;c=ws.cell(r,3).value;d=ws.cell(r,4)
    if b and str(b).strip():
        n=str(b).strip();cur=cats[0 if n.startswith('1.') else 1 if n.startswith('2.') else 2]['id'];continue
    if c and d.value:
        t=str(c).strip(); u=str(d.value).strip()
        if d.hyperlink and d.hyperlink.target and d.hyperlink.target.strip()!=u: diff.append((r,u,d.hyperlink.target))
        courses.append({'id':'c-%03d'%(len(courses)+1),'title':t,'categoryId':cur,'url':u})
print(len(courses),{k:sum(c['categoryId']==k for c in courses) for k in [x['id'] for x in cats]},'diff D value/hyperlink:',diff)
assert len({c['url'] for c in courses})==len(courses),'dup urls'
assert all(re.match(r'^https://my\.qgrnd\.kz/~\w+$',c['url']) for c in courses)
bk=json.load(open(U+'raspisanie-zoom-2026-09-30.json',encoding='utf-8'))
assert bk['app']=='zoom-schedule' and bk['version']==1
src=open(U+'index.html',encoding='utf-8').read()
demo=re.search(r'const DEMO = \[(.*?)\n\];',src,re.S)
print('demo ids in backup:',sum(b['id'].startswith('demo-') for b in bk['bookings']),'of',len(bk['bookings']))
cat={'schemaVersion':1,'revision':'2026-09-30-initial','exportedAt':'2026-09-30T00:00:00.000Z',
 'categories':cats,'courses':courses,
 'survey':{'title':'Оценка удовлетворенности слушателей','url':'https://docs.google.com/forms/d/e/1FAIpQLSftiDhA2JCpPOJ0dIRw_hN7geRymEkDOWM1pNKah1q2guQ7zQ/viewform?usp=dialog'},
 'schedule':{'resources':bk['resources'],'bookings':bk['bookings']}}
json.dump(cat,open('catalog.json','w',encoding='utf-8'),ensure_ascii=False,indent=1)
print([c['title'] for c in courses if '\n' in c['title'] or '  ' in c['title']])
